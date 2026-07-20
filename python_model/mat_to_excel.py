"""
NASA FDR — .mat to Excel Exporter
Converts one .mat file into a clean Excel workbook:
  Sheet 1 — Raw Data      (each parameter as a column, rows = time samples)
  Sheet 2 — Parameter Info (name, units, sample rate, description, stats)

Usage:
  python mat_to_excel.py
"""

import scipy.io
import numpy as np
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── CONFIG ────────────────────────────────────────────────────────────────────
DATA_DIR  = r"C:\Users\fatim\Desktop\Current pending papers\FDR phase 2\data\nasa_fdr_sample\Tail_652_7"
OUTPUT_DIR = r"C:\Users\fatim\Desktop\Research Assistant Work\ACoE\aviation-recorder-intelligence\python_model\fdr_segments"

# Pick which file to export (change index to try different files)
mat_files = sorted([f for f in os.listdir(DATA_DIR) if f.endswith('.mat')])
TARGET_FILE = mat_files[0]   # ← change to e.g. mat_files[5] for a different file

# Max rows to export in Raw Data sheet (Excel limit = 1,048,576)
# FDR files can be huge — cap at 50,000 rows for usability
MAX_ROWS = 50_000

os.makedirs(OUTPUT_DIR, exist_ok=True)

# ── LOAD ──────────────────────────────────────────────────────────────────────
fpath = os.path.join(DATA_DIR, TARGET_FILE)
print(f"📂 Loading: {TARGET_FILE}")
mat = scipy.io.loadmat(fpath, squeeze_me=False)
raw_params = {k: v for k, v in mat.items() if not k.startswith("__")}

# ── SMART EXTRACT ─────────────────────────────────────────────────────────────
def smart_extract(val):
    # Layout A: named struct fields
    try:
        inner = val[0][0]
        if inner.dtype.names:
            names = inner.dtype.names
            data  = inner[names[0]].flatten().astype(float)
            sr    = float(inner[names[1]].flatten()[0]) if len(names) > 1 else 0.0
            units = str(inner[names[2]]).strip()        if len(names) > 2 else "?"
            desc  = str(inner[names[3]]).strip()        if len(names) > 3 else "?"
            return data, sr, units, desc
    except: pass
    # Layout B: positional fields
    try:
        inner = val[0][0]
        data  = inner[0].flatten().astype(float)
        sr    = float(inner[1].flatten()[0]) if inner.shape[0] > 1 else 0.0
        units = str(inner[2]).strip()        if inner.shape[0] > 2 else "?"
        desc  = str(inner[3]).strip()        if inner.shape[0] > 3 else "?"
        return data, sr, units, desc
    except: pass
    # Layout C: plain array
    try:
        return val.flatten().astype(float), 0.0, "?", "?"
    except: pass
    raise ValueError("unreadable")

# ── PARSE ALL ─────────────────────────────────────────────────────────────────
print("🔄 Parsing parameters...")
parsed = {}
for key, val in raw_params.items():
    try:
        data, sr, units, desc = smart_extract(val)
        parsed[key] = {"data": data, "sr": sr, "units": units, "desc": desc}
    except Exception as e:
        print(f"  ⚠️  Skipping {key}: {e}")

print(f"  ✅ {len(parsed)} parameters parsed\n")

# ── FIND REFERENCE LENGTH (WOW or ALT) ───────────────────────────────────────
ref_key = next((k for k in ['WOW','ALT','CAS'] if k in parsed), list(parsed.keys())[0])
ref_len = len(parsed[ref_key]["data"])
print(f"  📌 Reference length: {ref_key} → {ref_len} samples")

# Resample every param to reference length
def resample(arr, target):
    if len(arr) == target:
        return arr
    idx = np.linspace(0, len(arr)-1, target).astype(int)
    return arr[idx]

# Cap rows
export_len = min(ref_len, MAX_ROWS)
if export_len < ref_len:
    print(f"  ✂️  Capping at {MAX_ROWS} rows (file has {ref_len} samples)")

# ── BUILD WORKBOOK ────────────────────────────────────────────────────────────
print("📊 Building Excel workbook...")
wb = Workbook()

# ── Styles ────────────────────────────────────────────────────────────────────
DARK    = "1F2937"
BLUE    = "2563EB"
GREEN   = "059669"
ORANGE  = "D97706"
RED     = "DC2626"
GRAY    = "F3F4F6"
WHITE   = "FFFFFF"

thin = Side(style='thin', color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def header_style(cell, bg=DARK, fg=WHITE, bold=True, center=True):
    cell.font = Font(name="Arial", bold=bold, color=fg, size=10)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center" if center else "left",
                                vertical="center", wrap_text=True)
    cell.border = border

def data_style(cell, bg=WHITE):
    cell.font = Font(name="Arial", size=9)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.border = border

# ── SHEET 1: RAW DATA ─────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Raw Data"
ws1.freeze_panes = "B2"

# Row 1: parameter keys
# Row 2: units
# Row 3 onwards: data

# Sort params: put key flight params first
PRIORITY = ['GMT_HOUR','GMT_MINUTE','GMT_SEC','WOW','ALT','RALT','CAS','TAS','MACH',
            'IVV','VRTG','PTCH','ROLL','LONG','LATG','AOA1','AOA2','FLAP',
            'N1_1','N1_2','N1_3','N1_4','N2_1','N2_2','EGT_1','EGT_2','FF_1','FF_2',
            'LATP','LONP']

ordered_keys = [k for k in PRIORITY if k in parsed]
ordered_keys += [k for k in parsed if k not in ordered_keys]

# Row index header
ws1.cell(row=1, column=1, value="Sample #")
header_style(ws1.cell(row=1, column=1), bg=DARK)
ws1.cell(row=2, column=1, value="—")
header_style(ws1.cell(row=2, column=1), bg="374151", fg="9CA3AF")
ws1.column_dimensions['A'].width = 10

for col_idx, key in enumerate(ordered_keys, start=2):
    p = parsed[key]
    col_letter = get_column_letter(col_idx)

    # Row 1: param name
    c1 = ws1.cell(row=1, column=col_idx, value=key)
    # Color code by category
    if key in ['ALT','RALT','CAS','TAS','IVV','VRTG','PTCH','ROLL','MACH','GS']:
        header_style(c1, bg=BLUE)
    elif any(key.startswith(p) for p in ['N1_','N2_','EGT_','FF_','PLA_','VIB_']):
        header_style(c1, bg="7C3AED")   # purple = engine
    elif key in ['WOW','FLAP','LGDN','LGUP','ABRK']:
        header_style(c1, bg=GREEN)
    elif key in ['LATP','LONP','GMT_HOUR','GMT_MINUTE','GMT_SEC','TRK','MH']:
        header_style(c1, bg=ORANGE)
    else:
        header_style(c1, bg="4B5563")

    # Row 2: units
    c2 = ws1.cell(row=2, column=col_idx, value=p["units"])
    header_style(c2, bg="374151", fg="9CA3AF", bold=False)

    # Column width
    ws1.column_dimensions[col_letter].width = max(len(key) + 2, 10)

print(f"  Writing {export_len} rows × {len(ordered_keys)} columns...")

# Write data rows
for row_idx in range(export_len):
    # Sample number
    ws1.cell(row=row_idx+3, column=1, value=row_idx+1)

    for col_idx, key in enumerate(ordered_keys, start=2):
        arr  = resample(parsed[key]["data"], ref_len)
        val  = arr[row_idx]
        cell = ws1.cell(row=row_idx+3, column=col_idx)

        if np.isnan(val):
            cell.value = None
        else:
            cell.value = round(float(val), 4)

        # Alternating row color
        bg = "F9FAFB" if row_idx % 2 == 0 else WHITE
        data_style(cell, bg=bg)

# ── SHEET 2: PARAMETER INFO ───────────────────────────────────────────────────
ws2 = wb.create_sheet("Parameter Info")
ws2.freeze_panes = "A2"

info_headers = ["#", "Parameter", "Description", "Units", "Sample Rate (Hz)",
                "N Samples", "Min", "Max", "Mean", "Std Dev", "NaN %",
                "Constant?", "Binary?", "Quality"]

col_widths = [5, 14, 45, 10, 18, 12, 12, 12, 12, 12, 8, 11, 10, 12]

for col_idx, (h, w) in enumerate(zip(info_headers, col_widths), start=1):
    c = ws2.cell(row=1, column=col_idx, value=h)
    header_style(c, bg=DARK)
    ws2.column_dimensions[get_column_letter(col_idx)].width = w

ws2.row_dimensions[1].height = 28

for row_idx, key in enumerate(ordered_keys, start=2):
    p   = parsed[key]
    arr = p["data"]
    valid = arr[~np.isnan(arr)]

    nan_pct  = round(np.sum(np.isnan(arr)) / len(arr) * 100, 1)
    is_const = len(valid) > 0 and float(np.std(valid)) < 1e-6
    is_bin   = len(valid) > 0 and len(np.unique(valid)) <= 2

    if nan_pct > 50:   quality = "❌ High NaN"
    elif is_const:     quality = "📌 Constant"
    elif is_bin:       quality = "🔘 Binary"
    elif nan_pct > 0:  quality = "⚠️ Has NaN"
    else:              quality = "✅ Clean"

    row_bg = "F9FAFB" if row_idx % 2 == 0 else WHITE

    values = [
        row_idx - 1,
        key,
        p["desc"][:60] if p["desc"] != "?" else "",
        p["units"] if p["units"] != "?" else "",
        p["sr"],
        len(arr),
        round(float(np.nanmin(arr)), 3) if len(valid) > 0 else "N/A",
        round(float(np.nanmax(arr)), 3) if len(valid) > 0 else "N/A",
        round(float(np.nanmean(arr)), 3) if len(valid) > 0 else "N/A",
        round(float(np.nanstd(arr)),  3) if len(valid) > 0 else "N/A",
        nan_pct,
        "Yes" if is_const else "No",
        "Yes" if is_bin   else "No",
        quality,
    ]

    for col_idx, val in enumerate(values, start=1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.font      = Font(name="Arial", size=9)
        cell.fill      = PatternFill("solid", start_color=row_bg)
        cell.border    = border
        cell.alignment = Alignment(
            horizontal="left" if col_idx in [2,3,4,14] else "center",
            vertical="center"
        )
        # Highlight quality cell
        if col_idx == 14:
            if "❌" in str(val): cell.fill = PatternFill("solid", start_color="FEE2E2")
            elif "⚠️" in str(val): cell.fill = PatternFill("solid", start_color="FEF9C3")
            elif "✅" in str(val): cell.fill = PatternFill("solid", start_color="DCFCE7")
            elif "📌" in str(val): cell.fill = PatternFill("solid", start_color="E5E7EB")
            elif "🔘" in str(val): cell.fill = PatternFill("solid", start_color="DBEAFE")

# ── SHEET 3: SUMMARY ─────────────────────────────────────────────────────────
ws3 = wb.create_sheet("Summary")

ok_params  = list(parsed.keys())
clean_ct   = sum(1 for k in ok_params if parsed[k]["data"][~np.isnan(parsed[k]["data"])].size > 0
                 and np.std(parsed[k]["data"][~np.isnan(parsed[k]["data"])]) >= 1e-6
                 and np.sum(np.isnan(parsed[k]["data"])) == 0)
nan_ct     = sum(1 for k in ok_params if 0 < np.sum(np.isnan(parsed[k]["data"])) / len(parsed[k]["data"]) <= 0.5)
highnan_ct = sum(1 for k in ok_params if np.sum(np.isnan(parsed[k]["data"])) / len(parsed[k]["data"]) > 0.5)
const_ct   = sum(1 for k in ok_params if len(parsed[k]["data"][~np.isnan(parsed[k]["data"])]) > 0
                 and np.std(parsed[k]["data"][~np.isnan(parsed[k]["data"])]) < 1e-6)

summary_data = [
    ("Source file",          TARGET_FILE),
    ("Total parameters",     len(parsed)),
    ("Reference signal",     ref_key),
    ("Total samples",        ref_len),
    ("Exported rows",        export_len),
    ("", ""),
    ("✅ Clean params",       clean_ct),
    ("⚠️  Has NaN (≤50%)",    nan_ct),
    ("❌ High NaN (>50%)",    highnan_ct),
    ("📌 Constant params",    const_ct),
    ("", ""),
    ("Sheet: Raw Data",      f"{export_len} rows × {len(ordered_keys)} columns"),
    ("Sheet: Parameter Info","One row per parameter with stats"),
    ("Sheet: Summary",       "This sheet"),
]

ws3.column_dimensions['A'].width = 28
ws3.column_dimensions['B'].width = 50

c = ws3.cell(row=1, column=1, value="FDR .mat → Excel Export Summary")
c.font = Font(name="Arial", bold=True, size=14, color=DARK)
c.alignment = Alignment(horizontal="left")

for i, (label, value) in enumerate(summary_data, start=3):
    ca = ws3.cell(row=i, column=1, value=label)
    cb = ws3.cell(row=i, column=2, value=value)
    if label:
        ca.font = Font(name="Arial", bold=True, size=10)
        cb.font = Font(name="Arial", size=10)
        bg = "F9FAFB" if i % 2 == 0 else WHITE
        ca.fill = cb.fill = PatternFill("solid", start_color=bg)
    ca.alignment = cb.alignment = Alignment(horizontal="left", vertical="center")

# ── SAVE ──────────────────────────────────────────────────────────────────────
out_name = TARGET_FILE.replace('.mat', '_data.xlsx')
out_path = os.path.join(OUTPUT_DIR, out_name)
wb.save(out_path)

print(f"\n✅ Done!")
print(f"   Saved to : {out_path}")
print(f"   Sheets   : Raw Data ({export_len} rows × {len(ordered_keys)} cols)  |  Parameter Info  |  Summary")
print(f"\n   💡 Tip: Open 'Parameter Info' sheet first to understand what each column means")
print(f"   💡 Tip: Change TARGET_FILE = mat_files[N] to export a different flight file")